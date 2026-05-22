from __future__ import annotations

import itertools
import json
import math
import random
import time
from dataclasses import dataclass
from pathlib import Path
from statistics import mean, pstdev
from typing import Dict, List, Sequence, Tuple

from openpyxl import load_workbook


DATA_PATH = Path(r"D:\Google_Install\expanded_milk_run_dataset.xlsx")
RESULT_PATH = Path("expanded_milk_run_results.json")

CAPACITY = 40.0
FREQUENCY = 4
RANDOM_SEED = 20260521


@dataclass(frozen=True)
class Supplier:
    code: str
    x: float
    y: float
    daily_total: float
    pickup: float
    type_count: int
    load_grade: str


def load_dataset(path: Path = DATA_PATH):
    wb = load_workbook(path, data_only=True)
    ws = wb["供应商数据"]
    suppliers: Dict[str, Supplier] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        code = str(row[0])
        suppliers[code] = Supplier(
            code=code,
            x=float(row[1]),
            y=float(row[2]),
            daily_total=float(row[9]),
            pickup=float(row[10]),
            type_count=int(row[11]),
            load_grade=str(row[12]),
        )

    ws_d = wb["距离矩阵"]
    headers = [str(v) for v in next(ws_d.iter_rows(min_row=1, max_row=1, values_only=True))[1:]]
    dist: Dict[str, Dict[str, float]] = {}
    for row in ws_d.iter_rows(min_row=2, values_only=True):
        src = str(row[0])
        dist[src] = {dst: float(value) for dst, value in zip(headers, row[1:])}

    return suppliers, dist


def route_load(route: Sequence[str], suppliers: Dict[str, Supplier]) -> float:
    return sum(suppliers[i].pickup for i in route)


def raw_route_distance(route: Sequence[str], dist: Dict[str, Dict[str, float]]) -> float:
    if not route:
        return 0.0
    total = dist["GT"][route[0]]
    for a, b in zip(route, route[1:]):
        total += dist[a][b]
    total += dist[route[-1]]["GT"]
    return total


def exact_route_order(route: Sequence[str], dist: Dict[str, Dict[str, float]]) -> Tuple[List[str], float]:
    route = list(route)
    if len(route) <= 1:
        return route, raw_route_distance(route, dist)
    best_order = None
    best_distance = float("inf")
    for perm in itertools.permutations(route):
        d = raw_route_distance(perm, dist)
        if d < best_distance:
            best_distance = d
            best_order = list(perm)
    return best_order or route, best_distance


def canonicalize_routes(routes: Sequence[Sequence[str]], dist, suppliers):
    improved = []
    for route in routes:
        ordered, _ = exact_route_order(route, dist)
        improved.append(ordered)
    return improved


def total_distance(routes: Sequence[Sequence[str]], dist) -> float:
    return sum(raw_route_distance(route, dist) for route in routes)


def metrics(routes: Sequence[Sequence[str]], dist, suppliers, runtime: float, complexity: str):
    loads = [route_load(r, suppliers) for r in routes]
    distances = [raw_route_distance(r, dist) for r in routes]
    load_rates = [v / CAPACITY for v in loads]
    return {
        "routes": [list(r) for r in routes],
        "route_count": len(routes),
        "single_total_distance": round(sum(distances), 2),
        "daily_total_distance": round(sum(distances) * FREQUENCY, 2),
        "single_total_load": round(sum(loads), 2),
        "avg_load_rate": round(mean(load_rates) * 100, 2),
        "min_load_rate": round(min(load_rates) * 100, 2),
        "max_load_rate": round(max(load_rates) * 100, 2),
        "load_rate_std": round(pstdev(load_rates) * 100, 2),
        "runtime_sec": round(runtime, 4),
        "complexity_note": complexity,
        "route_details": [
            {
                "route": list(r),
                "load": round(load, 2),
                "distance": round(distance, 2),
                "load_rate": round(rate * 100, 2),
            }
            for r, load, distance, rate in zip(routes, loads, distances, load_rates)
        ],
    }


def nearest_neighbor(suppliers, dist):
    start = time.perf_counter()
    unvisited = set(suppliers)
    routes = []
    while unvisited:
        route = []
        load = 0.0
        current = "GT"
        while True:
            feasible = [i for i in unvisited if load + suppliers[i].pickup <= CAPACITY + 1e-9]
            if not feasible:
                break
            nxt = min(feasible, key=lambda i: (dist[current][i], -suppliers[i].pickup, i))
            route.append(nxt)
            load += suppliers[nxt].pickup
            unvisited.remove(nxt)
            current = nxt
        routes.append(route)
    routes = canonicalize_routes(routes, dist, suppliers)
    return metrics(routes, dist, suppliers, time.perf_counter() - start, "贪心选择，约为O(n^2)")


def sweep_algorithm(suppliers, dist):
    start = time.perf_counter()
    items = list(suppliers)
    angles = {i: math.atan2(suppliers[i].y, suppliers[i].x) for i in items}
    ordered = sorted(items, key=lambda i: angles[i])
    best_routes = None
    best_distance = float("inf")
    for reverse in (False, True):
        seq0 = list(reversed(ordered)) if reverse else ordered
        for shift in range(len(seq0)):
            seq = seq0[shift:] + seq0[:shift]
            routes, route, load = [], [], 0.0
            for i in seq:
                demand = suppliers[i].pickup
                if route and load + demand > CAPACITY + 1e-9:
                    routes.append(route)
                    route, load = [], 0.0
                route.append(i)
                load += demand
            if route:
                routes.append(route)
            routes = canonicalize_routes(routes, dist, suppliers)
            d = total_distance(routes, dist)
            if d < best_distance:
                best_distance = d
                best_routes = routes
    return metrics(best_routes, dist, suppliers, time.perf_counter() - start, "极角排序与旋转搜索，约为O(n^2)")


def savings_algorithm(suppliers, dist):
    start = time.perf_counter()
    routes = {i: [i] for i in suppliers}
    route_of = {i: i for i in suppliers}
    loads = {i: suppliers[i].pickup for i in suppliers}
    savings = []
    ids = list(suppliers)
    for i, j in itertools.combinations(ids, 2):
        savings.append((dist["GT"][i] + dist["GT"][j] - dist[i][j], i, j))
    savings.sort(reverse=True)

    for _, i, j in savings:
        ri, rj = route_of[i], route_of[j]
        if ri == rj:
            continue
        route_i, route_j = routes[ri], routes[rj]
        if loads[ri] + loads[rj] > CAPACITY + 1e-9:
            continue
        candidates = []
        if route_i[-1] == i and route_j[0] == j:
            candidates.append(route_i + route_j)
        if route_i[0] == i and route_j[-1] == j:
            candidates.append(route_j + route_i)
        if route_i[0] == i and route_j[0] == j:
            candidates.append(list(reversed(route_i)) + route_j)
        if route_i[-1] == i and route_j[-1] == j:
            candidates.append(route_i + list(reversed(route_j)))
        if not candidates:
            continue
        merged = min(candidates, key=lambda r: raw_route_distance(r, dist))
        new_key = ri
        routes[new_key] = merged
        loads[new_key] = loads[ri] + loads[rj]
        del routes[rj]
        del loads[rj]
        for node in merged:
            route_of[node] = new_key

    out = canonicalize_routes(list(routes.values()), dist, suppliers)
    return metrics(out, dist, suppliers, time.perf_counter() - start, "节约值排序与端点合并，约为O(n^2 log n)")


def split_decode(perm: Sequence[str], suppliers, dist, route_penalty=18.0, balance_penalty=0.05):
    n = len(perm)
    best_cost = [float("inf")] * (n + 1)
    best_prev = [-1] * (n + 1)
    best_route = [None] * (n + 1)
    best_cost[0] = 0.0
    route_cache = {}
    for i in range(n):
        if best_cost[i] == float("inf"):
            continue
        load = 0.0
        subset = []
        for j in range(i, n):
            node = perm[j]
            load += suppliers[node].pickup
            if load > CAPACITY + 1e-9:
                break
            subset.append(node)
            key = tuple(sorted(subset))
            if key not in route_cache:
                route_cache[key] = exact_route_order(subset, dist)
            order, distance = route_cache[key]
            cost = (
                best_cost[i]
                + distance
                + route_penalty
                + balance_penalty * (CAPACITY - load) ** 2
            )
            if cost < best_cost[j + 1]:
                best_cost[j + 1] = cost
                best_prev[j + 1] = i
                best_route[j + 1] = list(order)
    routes = []
    idx = n
    while idx > 0:
        routes.append(best_route[idx])
        idx = best_prev[idx]
    routes.reverse()
    return routes


def ga_algorithm(suppliers, dist, population_size=120, generations=260, crossover_rate=0.86, mutation_rate=0.22, elite_size=6):
    start = time.perf_counter()
    rng = random.Random(RANDOM_SEED)
    ids = list(suppliers)
    seed_routes = [
        nearest_neighbor(suppliers, dist)["routes"],
        sweep_algorithm(suppliers, dist)["routes"],
        savings_algorithm(suppliers, dist)["routes"],
    ]

    def flatten(routes):
        return [node for route in routes for node in route]

    population = [flatten(r) for r in seed_routes]
    while len(population) < population_size:
        perm = ids[:]
        rng.shuffle(perm)
        population.append(perm)

    cache = {}

    def evaluate(perm):
        key = tuple(perm)
        if key in cache:
            return cache[key]
        routes = split_decode(perm, suppliers, dist)
        d = total_distance(routes, dist)
        loads = [route_load(r, suppliers) for r in routes]
        imbalance = pstdev([v / CAPACITY for v in loads]) if len(loads) > 1 else 0.0
        cost = d + 7.5 * len(routes) + 8.0 * imbalance
        cache[key] = (cost, routes)
        return cache[key]

    def tournament():
        candidates = rng.sample(population, 4)
        return min(candidates, key=lambda p: evaluate(p)[0])[:]

    def ordered_crossover(a, b):
        n = len(a)
        i, j = sorted(rng.sample(range(n), 2))
        child = [None] * n
        child[i:j] = a[i:j]
        fill = [x for x in b if x not in child]
        pos = 0
        for k in range(n):
            if child[k] is None:
                child[k] = fill[pos]
                pos += 1
        return child

    def mutate(p):
        q = p[:]
        if rng.random() < 0.45:
            i, j = rng.sample(range(len(q)), 2)
            q[i], q[j] = q[j], q[i]
        elif rng.random() < 0.75:
            i, j = rng.sample(range(len(q)), 2)
            node = q.pop(i)
            q.insert(j, node)
        else:
            i, j = sorted(rng.sample(range(len(q)), 2))
            q[i:j] = reversed(q[i:j])
        return q

    for _ in range(generations):
        ranked = sorted(population, key=lambda p: evaluate(p)[0])
        next_population = [p[:] for p in ranked[:elite_size]]
        while len(next_population) < population_size:
            p1, p2 = tournament(), tournament()
            if rng.random() < crossover_rate:
                child = ordered_crossover(p1, p2)
            else:
                child = p1[:]
            if rng.random() < mutation_rate:
                child = mutate(child)
            next_population.append(child)
        population = next_population

    best = min(population, key=lambda p: evaluate(p)[0])
    routes = evaluate(best)[1]
    result = metrics(routes, dist, suppliers, time.perf_counter() - start, "群体搜索与repair解码，约为O(G·P·n^2)")
    result["parameters"] = {
        "population_size": population_size,
        "generations": generations,
        "crossover_rate": crossover_rate,
        "mutation_rate": mutation_rate,
        "elite_size": elite_size,
        "random_seed": RANDOM_SEED,
    }
    return result


def simulated_annealing(suppliers, dist, iterations=9000, start_temp=80.0, cooling=0.9975):
    start = time.perf_counter()
    rng = random.Random(RANDOM_SEED + 7)
    current = [node for route in savings_algorithm(suppliers, dist)["routes"] for node in route]

    def evaluate(perm):
        routes = split_decode(perm, suppliers, dist)
        return total_distance(routes, dist) + 7.5 * len(routes), routes

    current_cost, current_routes = evaluate(current)
    best_perm, best_cost, best_routes = current[:], current_cost, current_routes
    temp = start_temp
    n = len(current)
    for _ in range(iterations):
        candidate = current[:]
        op = rng.random()
        if op < 0.45:
            i, j = rng.sample(range(n), 2)
            candidate[i], candidate[j] = candidate[j], candidate[i]
        elif op < 0.75:
            i, j = rng.sample(range(n), 2)
            node = candidate.pop(i)
            candidate.insert(j, node)
        else:
            i, j = sorted(rng.sample(range(n), 2))
            candidate[i:j] = reversed(candidate[i:j])
        candidate_cost, candidate_routes = evaluate(candidate)
        delta = candidate_cost - current_cost
        if delta < 0 or rng.random() < math.exp(-delta / max(temp, 1e-9)):
            current, current_cost, current_routes = candidate, candidate_cost, candidate_routes
            if current_cost < best_cost:
                best_perm, best_cost, best_routes = current[:], current_cost, current_routes
        temp *= cooling
    result = metrics(best_routes, dist, suppliers, time.perf_counter() - start, "邻域扰动与退火接受准则，约为O(K·n^2)")
    result["parameters"] = {
        "iterations": iterations,
        "start_temp": start_temp,
        "cooling": cooling,
        "random_seed": RANDOM_SEED + 7,
    }
    return result


def exact_partition_reference(suppliers, dist):
    """Small-scale set-partitioning reference used to validate the expanded case.

    This is not treated as the main solution method in the paper. It is included
    to distinguish the volume lower bound from an actually feasible fixed-route
    capacity plan.
    """
    start = time.perf_counter()
    ids = list(suppliers)
    index = {node: pos for pos, node in enumerate(ids)}
    feasible_by_first = {}
    max_route_size = len(ids)
    for size in range(1, max_route_size + 1):
        any_feasible = False
        for subset in itertools.combinations(ids, size):
            load = route_load(subset, suppliers)
            if load <= CAPACITY + 1e-9:
                any_feasible = True
                mask = sum(1 << index[node] for node in subset)
                order, distance = exact_route_order(subset, dist)
                first = min(index[node] for node in subset)
                feasible_by_first.setdefault(first, []).append((mask, distance, order))
        if not any_feasible and size > 4:
            break
    for entries in feasible_by_first.values():
        entries.sort(key=lambda item: item[1])

    all_mask = (1 << len(ids)) - 1
    memo = {}

    def search(mask):
        if mask == all_mask:
            return 0.0, []
        if mask in memo:
            return memo[mask]
        remaining = (~mask) & all_mask
        first = (remaining & -remaining).bit_length() - 1
        best_cost, best_routes = float("inf"), None
        for submask, distance, order in feasible_by_first[first]:
            if submask & mask:
                continue
            tail_cost, tail_routes = search(mask | submask)
            cost = distance + tail_cost
            if cost < best_cost:
                best_cost = cost
                best_routes = [order] + tail_routes
        memo[mask] = best_cost, best_routes
        return memo[mask]

    _, routes = search(0)
    result = metrics(routes, dist, suppliers, time.perf_counter() - start, "容量可行分组校验，精确集合划分搜索")
    return result


def main():
    suppliers, dist = load_dataset()
    algorithms = {
        "节点路径最短法": nearest_neighbor,
        "扫描法": sweep_algorithm,
        "节约里程法": savings_algorithm,
        "遗传算法": ga_algorithm,
        "模拟退火算法": simulated_annealing,
    }
    results = {}
    for name, fn in algorithms.items():
        print(f"running {name} ...")
        results[name] = fn(suppliers, dist)
    payload = {
        "data_path": str(DATA_PATH),
        "capacity": CAPACITY,
        "frequency": FREQUENCY,
        "supplier_count": len(suppliers),
        "daily_total_demand": round(sum(s.daily_total for s in suppliers.values()), 2),
        "single_total_demand": round(sum(s.pickup for s in suppliers.values()), 2),
        "theoretical_min_vehicles": math.ceil(sum(s.pickup for s in suppliers.values()) / CAPACITY),
        "suppliers": {k: suppliers[k].__dict__ for k in sorted(suppliers)},
        "results": results,
        "exact_partition_reference": exact_partition_reference(suppliers, dist),
    }
    RESULT_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    for name, r in results.items():
        print(
            name,
            "routes=", r["route_count"],
            "single_km=", r["single_total_distance"],
            "daily_km=", r["daily_total_distance"],
            "avg_load=", r["avg_load_rate"],
            "std=", r["load_rate_std"],
            "time=", r["runtime_sec"],
        )


if __name__ == "__main__":
    main()
